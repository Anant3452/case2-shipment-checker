# Builds the 8-slide submission deck. Run: python3 deck/make_deck.py
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

NAVY = RGBColor(0x1F, 0x4E, 0x8C)
INK = RGBColor(0x1D, 0x24, 0x30)
MUTED = RGBColor(0x5B, 0x65, 0x73)
LIGHT = RGBColor(0xF2, 0xF5, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIVE = "https://case2-shipment-checker.onrender.com"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, lines, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return box


def header(slide, n, title):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    text(slide, 0.6, 0.35, 12, 0.8, [title], size=30, color=NAVY, bold=True)
    text(slide, 0.6, 7.0, 10, 0.4, [f"Case 2 · {LIVE.replace('https://', '')}"], size=11, color=MUTED)
    text(slide, 12.0, 7.0, 0.8, 0.4, [str(n)], size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, title, body, fill=LIGHT):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill; shape.line.color.rgb = RGBColor(0xDF, 0xE3, 0xEA)
    shape.adjustments[0] = 0.06
    text(slide, x + 0.2, y + 0.15, w - 0.4, 0.5, [title], size=17, color=NAVY, bold=True)
    text(slide, x + 0.2, y + 0.65, w - 0.4, h - 0.8, body, size=14)


def node(slide, x, y, w, h, label, sub):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = NAVY; shape.line.fill.background()
    tf = shape.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = WHITE
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(11); r2.font.color.rgb = WHITE


def arrow(slide, x, y, w):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.35))
    a.fill.solid(); a.fill.fore_color.rgb = MUTED; a.line.fill.background()


# 1. Title ---------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
text(s, 0.8, 1.6, 11.5, 1.6, ["Catch shipment mistakes before they leave"], size=44, color=WHITE, bold=True)
text(s, 0.8, 3.1, 11.5, 0.8, ["Case 2: Conversational Shipment Booking & Pre-Flight Error Checker"], size=22, color=WHITE)
text(s, 0.8, 4.4, 11.5, 1.6, [
    "Team: Anant Jain and [partner name]",
    f"MVP link: {LIVE}",
    "Video link: [add before submitting]",
], size=18, color=WHITE)

# 2. Problem breakdown --------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, 2, "Problem breakdown")
card(s, 0.6, 1.4, 4.0, 5.3, "Who has it", [
    "UK sellers shipping stock to warehouses in Germany, France and the Netherlands.",
    "Often small teams with no freight expert in the room.",
])
card(s, 4.8, 1.4, 4.0, 5.3, "What it costs today", [
    "Goods held at customs for a wrong HS code or declared value.",
    "Units reported missing at the warehouse, because counts did not match.",
    "Found days later, when stock is already stuck.",
])
card(s, 9.0, 1.4, 3.8, 5.3, "Sub-problems it splits into", [
    "1. Forms are long and confusing.",
    "2. Numbers do not tally (value, cartons, units).",
    "3. HS codes are missing or too short.",
    "4. Carrier weight limits are not checked.",
    "5. Restricted goods are not flagged.",
])

# 3. Scope ------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, 3, "Scope: what we solved and what we left out")
card(s, 0.6, 1.4, 6.0, 5.3, "Solved in the MVP", [
    "Plain-language booking in chat, with only the missing questions asked.",
    "Upload of Bookings and Items sheets, with cell references for every problem.",
    "Rules R-01 to R-10, run the same way in both paths.",
    "Carrier ranking by seller priority, with excluded carriers and reasons.",
    "A booking summary the seller can copy.",
], fill=RGBColor(0xE6, 0xF4, 0xEC))
card(s, 6.9, 1.4, 5.9, 5.3, "Left out, and why", [
    "Real carrier bookings and labels: not needed to catch errors, and would need partner contracts.",
    "Customs filing and HS lookup: a legal service, not a checker.",
    "Payments, accounts and live exchange rates: GBP only, no login, for the MVP.",
    "Lanes beyond UK to DE, FR and NL: keeps the rate card honest.",
])

# 4. Assumptions & data ------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, 4, "Assumptions and how we would check them")
rows = [
    ("Rate card is accurate", "Only one carrier row is real; the rest is placeholder. Check against each carrier's price sheet before pilot."),
    ("1% value tolerance is right", "Measure how often honest sellers land between 0.5% and 1%. Tighten if rounding causes false alarms."),
    ("Balanced weights 0.4 / 0.3 / 0.3", "Ask 10 sellers to rank two options. Adjust weights to match what they choose."),
    ("Transit counted in calendar days", "Compare expected arrival with actual delivery dates for the pilot lane."),
    ("Keyword list catches restricted goods", "Review every R-06 hit and every miss in the pilot. Extend the list in one config file."),
]
y = 1.45
for a, b in rows:
    card(s, 0.6, y, 12.2, 1.0, a, [b])
    y += 1.1

# 5. Solution & MVP ----------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, 5, "Solution and MVP, end to end")
node(s, 0.5, 2.6, 2.3, 1.4, "Seller", "chat or sheets")
arrow(s, 2.9, 3.1, 0.6)
node(s, 3.6, 2.6, 2.5, 1.4, "Capture", "Gemini fills fields; fallback if no key")
arrow(s, 6.2, 3.1, 0.6)
node(s, 6.9, 2.6, 2.5, 1.4, "Checks", "R-01 to R-10 in code, not the model")
arrow(s, 9.5, 3.1, 0.6)
node(s, 10.2, 2.6, 2.6, 1.4, "Result", "status, carrier, summary")
text(s, 0.6, 4.6, 12.2, 2.0, [
    "The model only reads and writes fields. Pass or fail, carrier limits and totals come from deterministic code.",
    "Uploads go through the same checks. Each problem points at a spreadsheet cell.",
    f"Try it: {LIVE}",
], size=16)

# 6. Adoption ---------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, 6, "Adoption: who uses it, and why they might resist")
card(s, 0.6, 1.4, 4.0, 5.3, "Day-to-day user", [
    "Owner or operations lead at a small seller, 1 to 5 shipments a month.",
    "Books in the browser before the courier collects.",
])
card(s, 4.8, 1.4, 4.0, 5.3, "What changes for them", [
    "Replaces the carrier's long booking form with a short conversation.",
    "Errors are shown before collection, not days later at the warehouse.",
    "They keep their own sheets; upload works with what they already have.",
])
card(s, 9.0, 1.4, 3.8, 5.3, "Why they might resist", [
    "Distrust of AI giving wrong numbers. Mitigation: the model never decides pass or fail.",
    "Already have a broker. Mitigation: the checks are a second pair of eyes, not a replacement.",
    "Time to learn a new tool. Mitigation: one screen, one question at a time.",
])

# 7. Rollout ----------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, 7, "Rollout: pilot, phases, and the proof metrics")
card(s, 0.6, 1.4, 4.0, 5.3, "Pilot (weeks 1 to 4)", [
    "10 sellers on one lane: UK to Germany.",
    "Checks run alongside their usual bookings. Nothing is blocked yet.",
    "Every finding reviewed by us within 24 hours.",
])
card(s, 4.8, 1.4, 4.0, 5.3, "Phases", [
    "1. Shadow mode on one lane.",
    "2. Blocking mode on that lane, with a confirm-anyway option.",
    "3. Add France and the Netherlands, then import real rate cards.",
])
card(s, 9.0, 1.4, 3.8, 5.3, "Proof metrics", [
    "Share of bookings blocked before collection, and why.",
    "Held or short-counted shipments versus the pilot baseline.",
    "Repeat use: sellers booking 3 or more times a month.",
    "False alarms reported by sellers, target under 5%.",
])

# 8. Open slide -------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, 8, "Open slide: what we would build next")
card(s, 0.6, 1.4, 6.0, 5.3, "Brownie point: features next", [
    "Import real carrier rate cards and keep them current.",
    "Save bookings and re-check them when the ready date changes.",
    "Let sellers attach their own HS code lookups, reviewed by a broker.",
])
card(s, 6.9, 1.4, 5.9, 5.3, "Known limits to be honest about", [
    "Placeholder rate card until real data is loaded.",
    "The rule-based assistant handles one or two facts a message; the Gemini assistant handles longer messages.",
    "Free hosting sleeps when idle, so the first visit can be slow.",
])

prs.save("deck/Case2-shipment-checker-deck.pptx")
print("saved", len(prs.slides), "slides")
